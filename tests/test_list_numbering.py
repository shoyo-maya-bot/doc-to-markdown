"""Word の番号付き箇条書き: 段落や表を挟んでも続き番号になる（numId・段ごと、振り直し指定で戻す）."""

from __future__ import annotations

import shutil
import tempfile
import unittest
from pathlib import Path

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from doc_to_markdown.converter import convert_file
from doc_to_markdown.model import ListItem, Options, Paragraph
from doc_to_markdown.render import render


def _set_num(paragraph, num_id: str, ilvl: int = 0) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    ppr.append(parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/><w:numId w:val="{num_id}"/></w:numPr>'))


def _add_numbering(document, *, start: int = 1) -> tuple[str, str]:
    """10 進の番号付け（abstractNum）と、それを参照する num を 2 つ（2 つ目は 1 から振り直し）作る."""
    numbering = document.part.numbering_part.element
    abs_id, n1, n2 = "90", "91", "92"
    lvls = "".join(
        f'<w:lvl w:ilvl="{i}"><w:start w:val="{start}"/><w:numFmt w:val="decimal"/></w:lvl>' for i in range(2)
    )
    numbering.insert(0, parse_xml(f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{abs_id}">{lvls}</w:abstractNum>'))
    numbering.append(parse_xml(f'<w:num {nsdecls("w")} w:numId="{n1}"><w:abstractNumId w:val="{abs_id}"/></w:num>'))
    numbering.append(
        parse_xml(
            f'<w:num {nsdecls("w")} w:numId="{n2}"><w:abstractNumId w:val="{abs_id}"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="1"/></w:lvlOverride></w:num>'
        )
    )
    return n1, n2


class ListNumberingTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp)

    def convert(self, document) -> str:
        path = self.tmp / "list.docx"
        document.save(path)
        r = convert_file(path, self.tmp / "out", options=Options())
        return r.output.read_text(encoding="utf-8")

    def test_list_number_style_continues_across_paragraph(self):
        d = docx.Document()
        d.add_paragraph("手順 A", style="List Number")
        d.add_paragraph("手順 B", style="List Number")
        d.add_paragraph("補足の段落")
        d.add_paragraph("手順 C", style="List Number")
        md = self.convert(d)
        self.assertIn("1. 手順 A\n2. 手順 B", md)
        self.assertIn("3. 手順 C", md)
        self.assertNotIn("1. 手順 C", md)

    def test_num_id_continues_and_start_override_restarts(self):
        d = docx.Document()
        n1, n2 = _add_numbering(d)
        for text in ("一", "二"):
            _set_num(d.add_paragraph(text), n1)
        d.add_paragraph("間の段落")
        _set_num(d.add_paragraph("三"), n1)
        d.add_table(rows=1, cols=2).cell(0, 0).text = "表"
        _set_num(d.add_paragraph("四"), n1)
        _set_num(d.add_paragraph("四の一"), n1, ilvl=1)
        d.add_paragraph("段落")
        _set_num(d.add_paragraph("四の二"), n1, ilvl=1)
        _set_num(d.add_paragraph("五"), n1)
        _set_num(d.add_paragraph("五の一"), n1, ilvl=1)  # 上の段が進んだので 1 から
        d.add_paragraph("別の一覧")
        _set_num(d.add_paragraph("新しい一"), n2)  # startOverride: 振り直し
        md = self.convert(d)
        # 段落の直後の 2 段目は Markdown では入れ子にできないので 1 段目に出る（番号は続く）
        for line in (
            "1. 一",
            "2. 二",
            "3. 三",
            "4. 四",
            "  1. 四の一",
            "2. 四の二",
            "5. 五",
            "  1. 五の一",
            "1. 新しい一",
        ):
            self.assertIn(line + "\n", md)

    def test_start_value_is_used(self):
        d = docx.Document()
        n1, _ = _add_numbering(d, start=5)
        _set_num(d.add_paragraph("最初"), n1)
        d.add_paragraph("段落")
        _set_num(d.add_paragraph("次"), n1)
        md = self.convert(d)
        self.assertIn("5. 最初\n", md)
        self.assertIn("6. 次\n", md)

    def test_render_uses_given_number(self):
        md = render([ListItem("a", 0, True, number=3), Paragraph("p"), ListItem("b", 0, True, number=4)])
        self.assertIn("3. a", md)
        self.assertIn("4. b", md)

    def test_render_without_number_counts_from_one(self):
        md = render([ListItem("a", 0, True), Paragraph("p"), ListItem("b", 0, True)])
        self.assertIn("1. a", md)
        self.assertIn("1. b", md)


if __name__ == "__main__":
    unittest.main()
