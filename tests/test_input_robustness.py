"""壊れた入力の終了コード・脚注とコメントのエスケープ・出力先の除外・展開上限と外部コマンドの時間切れ."""

from __future__ import annotations

import base64
import json
import shutil
import subprocess
import urllib.parse
import zipfile
import zlib
from pathlib import Path
from unittest import mock

from devtools_common.cli import UsageError
from test_edge_cases import Base, run_cli

from doc_to_markdown import converter
from doc_to_markdown.adapters import pdf_adapter
from doc_to_markdown.converter import collect_inputs, convert_file
from doc_to_markdown.model import Document, FootnoteDef, Options
from doc_to_markdown.render import render

SAMPLES = Path(__file__).resolve().parents[1] / "samples"


# -- 1. 壊れた入力は入力の誤り（終了コード 2）----------------------------------
class BrokenInputExitCodeTest(Base):
    def _broken_files(self) -> list[Path]:
        files = []
        for ext in (".docx", ".xlsx", ".pptx", ".pdf", ".drawio"):
            p = self.tmp / f"text{ext}"  # 拡張子だけ Office・PDF のテキスト
            p.write_text("これはテキストです", encoding="utf-8")
            files.append(p)
        for name in ("expense-claim.docx", "expense-list.xlsx", "expense-flow.pptx", "expense-policy.pdf"):
            data = (SAMPLES / name).read_bytes()
            p = self.tmp / f"cut-{name}"  # 途中で切れたファイル
            p.write_bytes(data[: len(data) // 2])
            files.append(p)
        other = self.tmp / "other-zip.docx"  # Word の部品の無い zip
        with zipfile.ZipFile(other, "w") as z:
            z.writestr("readme.txt", "x")
        files.append(other)
        return files

    def test_each_broken_input_exits_2(self):
        for f in self._broken_files():
            with self.subTest(f.name):
                rc, _, err = run_cli(str(f), "-d", str(self.tmp / "out"))
                self.assertEqual(rc, 2, err)
                self.assertIn("読み取れません", err)
                self.assertNotIn("Traceback", err)
                self.assertFalse((self.tmp / "out" / f"{f.stem}.md").exists())

    def test_convert_file_raises_usage_error(self):
        f = self.tmp / "text.xlsx"
        f.write_text("x", encoding="utf-8")
        with self.assertRaises(UsageError):
            convert_file(f, self.tmp / "out")

    def test_continue_on_error_counts_broken_input_as_partial_failure(self):
        good = self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="a.xlsx")
        bad = self.tmp / "b.docx"
        bad.write_text("not a zip", encoding="utf-8")
        rc, out, _ = run_cli(str(good), str(bad), "-d", str(self.tmp / "out"), "--continue-on-error", "-f", "json")
        self.assertEqual(rc, 1)
        data = json.loads(out)
        errs = [f["location"] for f in data["findings"] if f["rule_id"] == "convert.input_error"]
        self.assertEqual(errs, ["b.docx"])
        self.assertEqual(data["meta"]["converted"], 1)

    def test_unexpected_bug_is_still_runtime_error(self):
        good = self.book([["区分", "金額"], ["A", 1]], name="a.xlsx")
        with mock.patch.dict(converter.ADAPTERS, {".xlsx": mock.Mock(side_effect=RuntimeError("bug"))}):
            rc, _, _ = run_cli(str(good), "-d", str(self.tmp / "out"))
        self.assertEqual(rc, 3)


# -- 2. 脚注・コメントのエスケープ ---------------------------------------------
class FootnoteEscapeTest(Base):
    def test_footnote_html_and_entities_are_escaped(self):
        md = render([FootnoteDef("1", "a <b>太字</b> と <!-- 消える --> と &lt;")])
        self.assertEqual(md, "[^1]: a \\<b>太字\\</b> と \\<!-- 消える --> と \\&lt;\n")

    def test_footnote_keeps_converter_markup(self):
        md = render([FootnoteDef("c1", "コメント（作成者）: **重要** <ins>追加</ins> \\*")])
        self.assertEqual(md, "[^c1]: コメント（作成者）: **重要** <ins>追加</ins> \\*\n")

    def test_docx_comment_with_html_is_escaped(self):
        import docx

        d = docx.Document()
        p = d.add_paragraph("本文")
        d.add_comment(p.runs, text="<script>x</script> と <!-- c -->", author="担当")
        src = self.tmp / "c.docx"
        d.save(str(src))
        r, md = self.convert(src)
        self.assertIn("\\<script>x\\</script> と \\<!-- c -->", md)
        self.assertNotIn(" <script>", md)
        self.assertNotIn("loss.text", self.rules(r))


# -- 3. 出力先が入力フォルダの中にあるとき ---------------------------------------
class OutputInsideInputTest(Base):
    def _docs(self) -> Path:
        docs = self.tmp / "docs"
        docs.mkdir()
        shutil.copy(SAMPLES / "approval-flow.drawio", docs / "flow.drawio")
        self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="docs/a.xlsx")
        return docs

    def test_rerun_does_not_collect_previous_output(self):
        docs = self._docs()
        out = docs / "md"
        for _ in range(2):
            rc, report, _ = run_cli(str(docs), "-r", "-d", str(out), "-f", "json")
            self.assertEqual(rc, 0)
            data = json.loads(report)
            self.assertEqual(data["meta"]["inputs"], 2)
        self.assertEqual(sorted(p.name for p in out.glob("*.md")), ["a.md", "flow.md"])

    def test_same_folder_skips_assets(self):
        docs = self._docs()
        rc, _, _ = run_cli(str(docs), "-r", "-d", str(docs))
        self.assertEqual(rc, 0)
        self.assertTrue((docs / "flow.assets" / "flow.drawio").is_file())  # draw.io のソースの写し
        files = collect_inputs([docs], recursive=True, out_dir=docs)
        self.assertEqual(sorted(f.name for f in files), ["a.xlsx", "flow.drawio"])

    def test_without_out_dir_collects_everything(self):
        docs = self._docs()
        (docs / "md" / "x.assets").mkdir(parents=True)
        shutil.copy(docs / "flow.drawio", docs / "md" / "x.assets" / "flow.drawio")
        self.assertEqual(len(collect_inputs([docs], recursive=True)), 3)
        self.assertEqual(len(collect_inputs([docs], recursive=True, out_dir=docs / "md")), 2)

    def test_explicit_file_inside_out_dir_is_kept(self):
        docs = self._docs()
        f = docs / "a.xlsx"
        self.assertEqual(collect_inputs([f], recursive=False, out_dir=docs), [f])


# -- 4. draw.io の展開上限（ファイル全体）と tesseract の時間切れ -----------------
class InflateTotalAndTimeoutTest(Base):
    @staticmethod
    def _page(xml: str, name: str) -> str:
        c = zlib.compressobj(9, zlib.DEFLATED, -15)
        raw = c.compress(urllib.parse.quote(xml).encode()) + c.flush()
        return f'<diagram name="{name}">{base64.b64encode(raw).decode()}</diagram>'

    def test_drawio_total_inflated_size_is_limited(self):
        pad = "<!--" + "a" * (600 << 10) + "-->"  # 1 ページは 1MB 未満、3 ページの合計は 1MB 超
        pages = "".join(self._page(f"<mxGraphModel>{pad}<root/></mxGraphModel>", f"p{i}") for i in range(3))
        path = self.tmp / "many.drawio"
        path.write_text(f"<mxfile>{pages}</mxfile>", encoding="utf-8")
        with self.assertRaises(UsageError) as cm:
            self.convert(path, max_unzipped_mb=1)
        self.assertIn("合計", str(cm.exception))
        r, _ = self.convert(path, max_unzipped_mb=4)
        self.assertEqual(r.source.name, "many.drawio")

    def test_ocr_page_timeout_is_reported(self):
        doc = Document()
        ocr = pdf_adapter._Ocr(Options(ocr="always"), doc)
        ocr.exe, ocr.lang = "tesseract", "jpn"
        page = mock.Mock()
        page.to_image.return_value.original.save = lambda p: Path(p).write_bytes(b"")
        with mock.patch.object(pdf_adapter.subprocess, "run", side_effect=subprocess.TimeoutExpired("t", 1)) as run:
            self.assertIsNone(ocr.read(page, 3))
        self.assertEqual(run.call_args.kwargs["timeout"], pdf_adapter.OCR_PAGE_TIMEOUT)
        ocr.finish()
        self.assertTrue(any("秒で終わらなかった" in w for w in doc.warnings))

    def test_ocr_list_langs_timeout(self):
        doc = Document()
        ocr = pdf_adapter._Ocr(Options(), doc)
        ocr.exe = "tesseract"
        with mock.patch.object(pdf_adapter.subprocess, "run", side_effect=subprocess.TimeoutExpired("t", 1)):
            self.assertIsNone(ocr.resolve_lang())

    def test_check_env_tesseract_has_timeout(self):
        from doc_to_markdown import cli

        def fake_find(spec):
            return "/x/tesseract" if spec.key == "tesseract" else None

        with (
            mock.patch("devtools_common.executables.find_executable", side_effect=fake_find),
            mock.patch("subprocess.run", side_effect=subprocess.TimeoutExpired("t", 1)) as run,
        ):
            rc, out, _ = run_cli("--check-env")
        self.assertEqual(rc, 0)
        self.assertEqual(run.call_args.kwargs["timeout"], cli.TESSERACT_TIMEOUT)
        self.assertIn("取得できませんでした", out)
