"""これまでに見つかった取りこぼし・誤変換の回帰テスト（意地悪な入力）.

- 年度・数値だけの見出し行、1 列だけの表、表を段落に落としたときの警告
- `[h]:mm` などの経過時間の表示形式
- 段落・見出しの Markdown 記法と HTML のエスケープ
- 大きなブックの省メモリ読み込みでも表示形式を反映し、warning で知らせる
- 出力の途中で失敗しても途中までの .md を残さない
- zip の展開サイズ・圧縮率の上限（zip 爆弾）
"""

from __future__ import annotations

import contextlib
import io
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest import mock

import openpyxl
from devtools_common.report import Severity
from openpyxl.styles import Font

from doc_to_markdown import cli, converter, lossaudit
from doc_to_markdown.adapters import drawio_adapter
from doc_to_markdown.converter import convert_file
from doc_to_markdown.model import Heading, Options, Paragraph
from doc_to_markdown.render import render


def run_cli(*args: str) -> tuple[int, str, str]:
    from devtools_common.cli import run_main

    out, err = io.StringIO(), io.StringIO()
    with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
        rc = run_main(cli.main, list(args))
    return rc, out.getvalue(), err.getvalue()


class Base(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        p = mock.patch.object(drawio_adapter, "find_drawio_cli", return_value=None)
        p.start()
        self.addCleanup(p.stop)

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def book(self, rows: list[list[object]], name: str = "book.xlsx", *, bold_rows=(), formats=None) -> Path:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for r, row in enumerate(rows, start=1):
            for c, v in enumerate(row, start=1):
                if v is None:
                    continue
                cell = ws.cell(r, c, v)
                if r in bold_rows:
                    cell.font = Font(bold=True)
        for ref, fmt in (formats or {}).items():
            ws[ref].number_format = fmt
        path = self.tmp / name
        wb.save(path)
        return path

    def convert(self, path: Path, **opts):
        r = convert_file(path, self.tmp / "out", options=Options(**opts))
        return r, r.output.read_text(encoding="utf-8")

    @staticmethod
    def rules(r, severity=None) -> list[str]:
        return [f.rule_id for f in r.findings if severity is None or f.severity is severity]


class YearHeaderTest(Base):
    """見出しが `項目 / 2024 / 2025 / 2026` の表（実務で見かける入力）."""

    ROWS = [
        ["項目", 2024, 2025, 2026],
        ["売上", 100, 120, 150],
        ["費用", 80, 90, 100],
    ]

    def test_year_header_is_table(self):
        r, md = self.convert(self.book(self.ROWS))
        self.assertIn("| 項目 | 2024 | 2025 | 2026 |\n| --- | --- | --- | --- |\n| 売上 | 100 | 120 | 150 |", md)
        self.assertNotIn("loss.tables", self.rules(r))
        self.assertNotIn("loss.table_as_paragraph", self.rules(r))

    def test_year_header_with_title_above(self):
        r, md = self.convert(self.book([["年度別の推移"], [], *self.ROWS]))
        self.assertIn("年度別の推移", md)
        self.assertIn("| 項目 | 2024 | 2025 | 2026 |", md)

    def test_fiscal_years_as_text(self):
        rows = [["項目", "2024", "2025"], ["売上", 1, 2], ["費用", 3, 4]]
        _, md = self.convert(self.book(rows))
        self.assertIn("| 項目 | 2024 | 2025 |\n| --- |", md)

    def test_numeric_header_distinguished_by_bold(self):
        # 年度でない数値の見出し（等級 1 / 2 / 3）も、見出しだけ太字なら見出しにする
        rows = [["等級", 1, 2, 3], ["A", 10, 20, 30], ["B", 40, 50, 60]]
        _, md = self.convert(self.book(rows, bold_rows=(1,)))
        self.assertIn("| 等級 | 1 | 2 | 3 |\n| --- |", md)

    def test_all_year_header_without_label(self):
        rows = [[2024, 2025, 2026], [100, 120, 150], [80, 90, 100]]
        _, md = self.convert(self.book(rows))
        self.assertIn("| 2024 | 2025 | 2026 |\n| --- | --- | --- |\n| 100 | 120 | 150 |", md)

    def test_streaming_mode_detects_year_header(self):
        rows = [["メモ"], ["項目", 2024, 2025], ["売上", 1, 2]]
        _, md = self.convert(self.book(rows), xlsx_full_load_max_cells=1)
        self.assertIn("メモ\n\n| 項目 | 2024 | 2025 |", md)


class SingleColumnTableTest(Base):
    def test_single_column_with_bold_header(self):
        rows = [["社員番号"], ["E001"], ["E002"], ["E003"]]
        r, md = self.convert(self.book(rows, bold_rows=(1,)))
        self.assertIn("| 社員番号 |\n| --- |\n| E001 |\n| E002 |\n| E003 |", md)
        self.assertNotIn("loss.table_as_paragraph", self.rules(r))

    def test_single_column_numbers_under_text_header(self):
        rows = [["数量"], [10], [20], [30]]
        _, md = self.convert(self.book(rows))
        self.assertIn("| 数量 |\n| --- |\n| 10 |", md)

    def test_title_and_notes_stay_paragraphs(self):
        rows = [["経費一覧"], ["作成: 総務部"], ["※ 4 月分"]]
        r, md = self.convert(self.book(rows, bold_rows=(1,)))
        self.assertNotIn("| --- |", md)
        self.assertIn("経費一覧\n\n作成: 総務部", md)
        self.assertNotIn("loss.table_as_paragraph", self.rules(r))

    def test_table_without_header_warns(self):
        # 見出しが無い（全部数値）の表は段落になるが、専用の警告を出す
        rows = [[1.5, 2.5, 3.5], [4.5, 5.5, 6.5], [7.5, 8.5, 9.5]]
        r, md = self.convert(self.book(rows))
        self.assertIn("1.5 / 2.5 / 3.5", md)
        found = [f for f in r.findings if f.rule_id == "loss.table_as_paragraph"]
        self.assertEqual(len(found), 1)
        self.assertIs(found[0].severity, Severity.WARNING)
        self.assertIn("Sheet1!A1:C3", found[0].message)
        self.assertIn("--header-row", found[0].suggestion)


class ElapsedTimeFormatTest(Base):
    def test_hours_over_24(self):
        path = self.book(
            [["作業", "時間", "秒まで", "分"], ["設計", 1.5, 1.5, 0.75]],
            formats={"B2": "[h]:mm", "C2": "[h]:mm:ss", "D2": "[mm]:ss"},
        )
        r, md = self.convert(path)
        self.assertIn("| 設計 | 36:00 | 36:00:00 | 1080:00 |", md)
        self.assertNotIn("day", md)
        self.assertNotIn("loss.number_format", self.rules(r))

    def test_streaming_mode_elapsed(self):
        path = self.book([["作業", "時間"], ["設計", 1.5]], formats={"B2": "[h]:mm"})
        _, md = self.convert(path, xlsx_full_load_max_cells=1)
        self.assertIn("| 設計 | 36:00 |", md)


class MarkdownEscapeTest(Base):
    """セル値・段落が Markdown 記法や HTML として解釈されないこと."""

    NASTY = [
        "# 大見出し",
        "## 小見出し",
        "- 箇条書きではない",
        "+ これも",
        "* これも",
        "1. 番号付きではない",
        "2) これも",
        "> 引用ではない",
        "<!-- 消えてはいけない -->",
        "<b>タグ</b>",
        "a | b | c",
        "C:\\temp\\file",
        "---",
        "===",
        "```",
        "[1]: http://example.invalid",
        "&lt;そのまま&gt;",
    ]

    def test_render_escapes_block_syntax(self):
        md = render([Paragraph(t) for t in self.NASTY])
        lines = md.split("\n\n")
        for src, line in zip(self.NASTY, lines, strict=True):
            with self.subTest(src=src):
                self.assertFalse(line.startswith(("#", "-", "+", "* ", ">", "<", "`", "=", "[")), line)
                self.assertNotRegex(line, r"^\d+[.)] ")
        self.assertIn("\\# 大見出し", md)
        self.assertIn("1\\. 番号付きではない", md)
        self.assertIn("\\<!-- 消えてはいけない -->", md)
        self.assertIn("a \\| b \\| c", md)
        self.assertIn("C:\\\\temp\\\\file", md)
        self.assertIn("\\&lt;そのまま\\&gt;", md)

    def test_every_line_of_paragraph_is_escaped(self):
        md = render([Paragraph("前置き\n# 見出しにしない\n- 箇条書きにしない\n===")])
        self.assertEqual(md, "前置き\\\n\\# 見出しにしない\\\n\\- 箇条書きにしない\\\n\\===\n")

    def test_heading_text_is_escaped(self):
        md = render([Heading(2, "<!-- x --> 見出し #")])
        self.assertEqual(md, "## \\<!-- x --> 見出し \\#\n")

    def test_markup_paragraph_keeps_inline_markup(self):
        # Word 由来の段落（太字・リンク・変更履歴）は行頭だけエスケープし、意図した記法は残す
        md = render([Paragraph("# **太字** [リンク](http://example.invalid) <ins>追加</ins>", markup=True)])
        self.assertEqual(md, "\\# **太字** [リンク](http://example.invalid) <ins>追加</ins>\n")

    def test_plain_text_counts_escaped_chars(self):
        md = render([Paragraph(t) for t in self.NASTY])
        self.assertEqual(
            lossaudit.markdown_stats(md).text_chars, lossaudit.count_chars("".join(self.NASTY).replace("|", ""))
        )

    def test_xlsx_cells_with_markdown_syntax(self):
        rows = [["# 大見出し"], ["<!-- 承認 -->"], ["1. 手順"]]
        r, md = self.convert(self.book(rows))
        self.assertIn("\\# 大見出し", md)
        self.assertIn("\\<!-- 承認 -->", md)
        self.assertIn("1\\. 手順", md)
        self.assertNotIn("loss.text", self.rules(r))
        self.assertNotIn("loss.raw_text", self.rules(r))

    def test_table_cell_html_is_escaped(self):
        rows = [["区分", "備考"], ["A", "<!-- 消えない -->"], ["B", "x"]]
        _, md = self.convert(self.book(rows))
        self.assertIn("| A | \\<!-- 消えない --> |", md)


class LargeWorkbookTest(Base):
    def test_streaming_keeps_number_format_and_warns(self):
        path = self.book([["コード", "名称"], [1, "本社"], [12, "支社"]], formats={"A2": "000", "A3": "000"})
        r, md = self.convert(path, xlsx_full_load_max_cells=3)
        self.assertIn("| 001 | 本社 |", md)
        self.assertIn("| 012 | 支社 |", md)
        found = [f for f in r.findings if f.rule_id == "convert.large_workbook"]
        self.assertEqual(len(found), 1)
        self.assertIs(found[0].severity, Severity.WARNING)
        self.assertIn("--max-cells", found[0].suggestion)

    def test_max_cells_option(self):
        path = self.book([["コード", "名称"], [1, "本社"]])
        rc, _, _ = run_cli(str(path), "-d", str(self.tmp / "o"), "--max-cells", "2", "--strict")
        self.assertEqual(rc, 1)  # 省メモリ読み込みは warning
        rc, _, _ = run_cli(str(path), "-d", str(self.tmp / "o2"), "--max-cells", "100", "--strict")
        self.assertEqual(rc, 0)


class AtomicWriteTest(Base):
    def test_failure_leaves_no_partial_output(self):
        path = self.book([["コード", "名称"], ["A", "本社"]])
        out = self.tmp / "out"
        out.mkdir()
        (out / "book.md").write_text("前回の出力\n", encoding="utf-8")

        def boom(blocks):
            yield from blocks
            raise RuntimeError("途中で失敗")

        with (
            mock.patch.object(converter, "normalize", side_effect=lambda b: boom(list(b))),
            self.assertRaises(RuntimeError),
        ):
            convert_file(path, out)
        self.assertEqual((out / "book.md").read_text(encoding="utf-8"), "前回の出力\n")  # 前回の出力は壊さない
        self.assertEqual(sorted(p.name for p in out.iterdir()), ["book.md"])  # 一時ファイルも残さない


class ZipBombTest(Base):
    def _bomb(self, name: str, size: int) -> Path:
        src = self.book([["a", "b"], ["1", "2"]], name="src.xlsx")
        path = self.tmp / name
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                zout.writestr(info, zin.read(info))
            zout.writestr("xl/media/padding.bin", b"\0" * size)
        return path

    def test_total_size_limit(self):
        path = self._bomb("big.xlsx", 3 * 1024 * 1024)
        with self.assertRaises(converter.UsageError) as cm:
            convert_file(path, self.tmp / "out", options=Options(max_unzipped_mb=2))
        self.assertIn("展開後", str(cm.exception))
        self.assertFalse((self.tmp / "out" / "big.md").exists())

    def test_ratio_limit(self):
        path = self._bomb("ratio.xlsx", 20 * 1024 * 1024)  # ゼロ埋め 20MB は 1000 倍以上に縮む
        with self.assertRaises(converter.UsageError) as cm:
            convert_file(path, self.tmp / "out")
        self.assertIn("圧縮率", str(cm.exception))

    def test_cli_exit_code(self):
        path = self._bomb("big.xlsx", 3 * 1024 * 1024)
        rc, _, err = run_cli(str(path), "-d", str(self.tmp / "out"), "--max-unzipped-mb", "2")
        self.assertEqual(rc, 2)
        self.assertIn("展開後", err)

    def test_docx_is_checked_too(self):
        src = Path(__file__).resolve().parent.parent / "samples" / "expense-claim.docx"
        path = self.tmp / "claim.docx"
        with zipfile.ZipFile(src) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                zout.writestr(info, zin.read(info))
            zout.writestr("word/media/padding.bin", b"\0" * (3 * 1024 * 1024))
        with self.assertRaises(converter.UsageError):
            convert_file(path, self.tmp / "out", options=Options(max_unzipped_mb=2))

    def test_normal_samples_pass(self):
        r, _ = self.convert(self.book([["a", "b"], ["1", "2"]]))
        self.assertTrue(r.output.is_file())


if __name__ == "__main__":
    unittest.main()


class BatchErrorsEscapeAndNotesTest(Base):
    """第 2 回修正: 続行モードでの使い方の誤り・各ブロックのエスケープ・注記行の警告・抽出図の保全・非表示行."""

    def test_continue_on_error_stops_on_usage_error(self):
        # 未対応の形式を直接指定したのは使い方の誤り. --continue-on-error でも続行せず、何も書かずに終了コード 2
        good = self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="good.xlsx")
        bad = self.tmp / "memo.txt"
        bad.write_text("x", encoding="utf-8")
        out_dir = self.tmp / "out"
        rc, _, err = run_cli(str(good), str(bad), "-d", str(out_dir), "--continue-on-error")
        self.assertEqual(rc, 2)
        self.assertIn("未対応", err)
        self.assertFalse((out_dir / "good.md").exists())

    def test_zip_limit_stops_without_continue_on_error(self):
        # --continue-on-error なしでは、上限超過（UsageError）で止まり終了コード 2
        good = self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="a.xlsx")
        big = self.tmp / "b.xlsx"
        with zipfile.ZipFile(big, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/big.xml", "0" * (2 << 20))
        rc, _, err = run_cli(str(good), str(big), "-d", str(self.tmp / "out"), "--max-unzipped-mb", "1")
        self.assertEqual(rc, 2)
        self.assertIn("上限", err)

    def test_list_item_is_escaped(self):
        from doc_to_markdown.model import ListItem

        md = render([ListItem("# 見出しではない"), ListItem("<!-- 消えない --> a|b", 1)])
        self.assertEqual(md, "- \\# 見出しではない\n  - \\<!-- 消えない --> a\\|b\n")

    def test_markup_list_item_keeps_inline_markup(self):
        from doc_to_markdown.model import ListItem

        md = render([ListItem("**太字** <ins>追加</ins>", markup=True)])
        self.assertEqual(md, "- **太字** <ins>追加</ins>\n")

    def test_quote_is_escaped(self):
        from doc_to_markdown.model import Quote

        md = render([Quote("# 図形の文字\n<!-- 消えない -->\n- 箇条書きではない", "SmartArt")])
        self.assertEqual(md, "> **SmartArt**: \\# 図形の文字\n> \\<!-- 消えない -->\n> \\- 箇条書きではない\n")

    def test_comment_cannot_be_closed_early(self):
        from doc_to_markdown.model import Comment

        md = render([Comment("a --> b\nc")])
        self.assertEqual(md.count("-->"), 1)
        self.assertTrue(md.startswith("<!-- a "))
        self.assertTrue(md.rstrip("\n").endswith(" -->"))
        self.assertNotIn("\n", md.rstrip("\n"))

    def test_trailing_note_rows_are_reported(self):
        rows = [["区分", "金額"], ["A", 1], ["B", 2], ["※ 税込"], ["注 端数切り捨て"]]
        r, md = self.convert(self.book(rows))
        self.assertIn("| B | 2 |", md)
        self.assertIn("※ 税込", md)
        warns = [f.message for f in r.findings if f.rule_id == "convert.warning"]
        self.assertTrue(any("注記らしい行 2 行" in w and "Sheet1!A4" in w and "Sheet1!A5" in w for w in warns), warns)

    def test_no_note_warning_without_note_rows(self):
        r, _ = self.convert(self.book([["区分", "金額"], ["A", 1], ["B", 2]]))
        self.assertFalse(any("注記らしい行" in f.message for f in r.findings))

    def test_assets_kept_when_conversion_fails(self):
        out_dir = self.tmp / "out"
        assets = out_dir / "book.assets"
        assets.mkdir(parents=True)
        (assets / "book-01.png").write_bytes(b"old")
        path = self.book([["区分", "金額"], ["A", 1], ["B", 2]])
        with (
            mock.patch.dict(converter.ADAPTERS, {".xlsx": mock.Mock(side_effect=RuntimeError("boom"))}),
            self.assertRaises(RuntimeError),
        ):
            convert_file(path, out_dir)
        self.assertEqual((assets / "book-01.png").read_bytes(), b"old")
        self.assertEqual(sorted(p.name for p in out_dir.iterdir()), ["book.assets"])  # 退避用の一時フォルダを残さない

    def test_stale_assets_removed_after_success(self):
        out_dir = self.tmp / "out"
        assets = out_dir / "book.assets"
        assets.mkdir(parents=True)
        (assets / "book-01.png").write_bytes(b"old")
        convert_file(self.book([["区分", "金額"], ["A", 1], ["B", 2]]), out_dir)
        self.assertFalse(assets.exists())
        self.assertEqual(sorted(p.name for p in out_dir.iterdir()), ["book.md"])

    def test_hidden_rows_counted_only_when_they_have_values(self):
        path = self.book([["区分", "金額"], ["A", 1], ["B", 2], ["C", 3]])
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        ws.row_dimensions[3].hidden = True
        ws.row_dimensions[10].hidden = True  # 値の無い非表示行は数えない
        wb.save(path)
        r, _ = self.convert(path)
        self.assertTrue(any("非表示の行 1 行" in f.message for f in r.findings))


class LinkEncodingAndPartialFailureTest(Base):
    """第 3 回修正: リンク先のエンコード・markup 段落のエスケープ・続行モードの入力誤り・draw.io の展開上限."""

    @staticmethod
    def _docx(path: Path, build) -> Path:
        import docx

        d = docx.Document()
        build(d)
        d.save(str(path))
        return path

    @staticmethod
    def _hyperlink(paragraph, url: str, text: str) -> None:
        from docx.opc.constants import RELATIONSHIP_TYPE as RT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        rid = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = text
        r.append(t)
        h.append(r)
        paragraph._p.append(h)

    # -- 1. リンク先 -----------------------------------------------------------
    def test_link_destination_encodes_space_and_parens(self):
        from doc_to_markdown.render import link_destination

        self.assertEqual(
            link_destination("a b (1).assets/a b (1)-01.png"), "a%20b%20%281%29.assets/a%20b%20%281%29-01.png"
        )
        self.assertEqual(link_destination("100%.assets/x.png"), "100%25.assets/x.png")
        self.assertEqual(link_destination("仕様書.assets/仕様書-01.png"), "仕様書.assets/仕様書-01.png")
        # URL の既にエンコード済みの部分は二重にしない
        self.assertEqual(
            link_destination("https://example.invalid/a%20b (c)", path=False), "https://example.invalid/a%20b%20%28c%29"
        )

    def test_image_block_path_with_space_and_parens(self):
        from doc_to_markdown.model import Image

        md = render([Image("図 (1).assets/図 (1)-01.png", alt="図")])
        self.assertEqual(md, "![図](図%20%281%29.assets/図%20%281%29-01.png)\n")

    def test_docx_image_and_hyperlink_with_space_and_parens(self):
        import base64

        png = base64.b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="
        )
        img = self.tmp / "p.png"
        img.write_bytes(png)

        def build(d):
            d.add_picture(str(img))
            p = d.add_paragraph("参照: ")
            self._hyperlink(p, "https://example.invalid/a b(c).html", "資料")
            t = d.add_table(rows=1, cols=1)
            t.cell(0, 0).paragraphs[0].add_run().add_picture(str(img))

        src = self._docx(self.tmp / "仕様 (改).docx", build)
        _, md = self.convert(src)
        self.assertIn("](仕様%20%28改%29.assets/仕様%20%28改%29-01.png)", md)
        self.assertNotIn("](仕様 (改)", md)
        self.assertIn("[資料](https://example.invalid/a%20b%28c%29.html)", md)
        # 表のセルの中の画像も同じ
        self.assertRegex(md, r"\| !\[[^\]]*\]\(仕様%20%28改%29\.assets/")

    def test_drawio_source_link_with_space_and_parens(self):
        src = self.tmp / "承認 フロー (案).drawio"
        shutil.copy(Path(__file__).resolve().parents[1] / "samples" / "approval-flow.drawio", src)
        _, md = self.convert(src)
        self.assertIn("[承認 フロー (案).drawio](承認%20フロー%20%28案%29.assets/承認%20フロー%20%28案%29.drawio)", md)

    def test_llms_txt_link_with_space(self):
        from doc_to_markdown import llmstxt

        out = self.tmp / "o"
        out.mkdir()
        md = out / "a b (1).md"
        md.write_text("# t\n", encoding="utf-8")
        llmstxt.write(out, [("a b (1).docx", md)])
        self.assertIn("(a%20b%20%281%29.md)", (out / "llms.txt").read_text(encoding="utf-8"))

    def test_llms_txt_link_text_escapes_brackets(self):
        # 再現（第 6 回）: 見出し「a]b(c」をそのまま [..] に入れ、`]` でリンクの文字が切れ、
        # `](` でリンク先が変わっていた
        from doc_to_markdown import llmstxt

        out = self.tmp / "o"
        out.mkdir()
        (out / "a]b(c.md").write_text("# a]b(c\n", encoding="utf-8")
        (out / "b.md").write_text("# x \\[y\\](z)\n", encoding="utf-8")  # 描画でエスケープ済みの [ ] はそのまま
        (out / "c.md").write_text("見出しなし\n", encoding="utf-8")
        items = [("a]b(c.xlsx", out / "a]b(c.md"), ("b-->.docx", out / "b.md"), ("c[1](2).xlsx", out / "c.md")]
        llmstxt.write(out, items)
        index = (out / "llms.txt").read_text(encoding="utf-8")
        self.assertIn("- [a\\]b(c](a]b%28c.md): 元ファイル a\\]b(c.xlsx", index)
        self.assertIn("- [x \\[y\\](z)](b.md):", index)
        self.assertIn("- [c](c.md): 元ファイル c\\[1\\](2).xlsx", index)
        # llms-full.txt のコメントは -- で閉じない
        self.assertIn("<!-- source: b- ->.docx -->", (out / "llms-full.txt").read_text(encoding="utf-8"))

    # -- 2. markup=True の段落のエスケープ --------------------------------------
    def test_docx_plain_text_symbols_are_escaped_but_bold_kept(self):
        def build(d):
            p = d.add_paragraph("パス C:\\temp\\*.txt と a_b_c と [注] と ~~x~~ と `code` ")
            p.add_run("太字").bold = True
            d.add_heading("見出し *重要*", level=1)
            d.add_paragraph("項目 *1*", style="List Bullet")

        r, md = self.convert(self._docx(self.tmp / "esc.docx", build))
        self.assertIn("パス C:\\\\temp\\\\\\*.txt と a\\_b\\_c と \\[注\\] と \\~\\~x\\~\\~ と \\`code\\` **太字**", md)
        self.assertIn("# 見出し \\*重要\\*", md)
        self.assertIn("- 項目 \\*1\\*", md)
        # エスケープの `\` は本文の文字数に入れない（ロス検知が誤検知しない・取りこぼしを見逃さない）
        self.assertNotIn("loss.text", self.rules(r))
        self.assertNotIn("loss.raw_text", self.rules(r))

    def test_docx_table_cell_symbols_are_escaped(self):
        def build(d):
            t = d.add_table(rows=2, cols=2)
            t.cell(0, 0).text = "区分"
            t.cell(0, 1).text = "パス"
            t.cell(1, 0).text = "*注*"
            t.cell(1, 1).paragraphs[0].add_run("C:\\a|b").bold = True

        _, md = self.convert(self._docx(self.tmp / "tbl.docx", build))
        self.assertIn("| \\*注\\* | **C:\\\\a\\|b** |", md)

    # -- 3. --continue-on-error でのファイル単位の入力の誤り ----------------------
    def _zip_bomb(self, name: str) -> Path:
        big = self.tmp / name
        with zipfile.ZipFile(big, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("xl/big.xml", "0" * (2 << 20))
        return big

    def test_continue_on_error_continues_on_zip_limit(self):
        import json

        good = self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="a.xlsx")
        big = self._zip_bomb("b.xlsx")
        good2 = self.book([["区分", "金額"], ["C", 3], ["D", 4]], name="c.xlsx")
        out_dir = self.tmp / "out"
        rc, out, _ = run_cli(
            str(good), str(big), str(good2), "-d", str(out_dir), "--continue-on-error",
            "--max-unzipped-mb", "1", "--llms-txt", "-f", "json",
        )  # fmt: skip
        self.assertEqual(rc, 1)  # 一部失敗（入力の誤り）は終了コード 1. 引数の誤り（2）とは区別する
        data = json.loads(out)
        self.assertEqual(data["meta"]["converted"], 2)
        errs = [f for f in data["findings"] if f["rule_id"] == "convert.input_error"]
        self.assertEqual([f["location"] for f in errs], ["b.xlsx"])
        self.assertIn("上限", errs[0]["message"])
        self.assertTrue((out_dir / "a.md").is_file())
        self.assertTrue((out_dir / "c.md").is_file())
        index = (out_dir / "llms.txt").read_text(encoding="utf-8")
        self.assertIn("(a.md)", index)
        self.assertIn("(c.md)", index)

    def test_continue_on_error_continues_without_libreoffice(self):
        import json

        good = self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="a.xlsx")
        old = self.tmp / "old.doc"
        old.write_bytes(b"\xd0\xcf\x11\xe0")
        out_dir = self.tmp / "out"
        with mock.patch("doc_to_markdown.office.find_soffice", return_value=None):
            rc, out, _ = run_cli(str(good), str(old), "-d", str(out_dir), "--continue-on-error", "-f", "json")
        self.assertEqual(rc, 1)
        data = json.loads(out)
        self.assertEqual(data["meta"]["converted"], 1)
        errs = [f["location"] for f in data["findings"] if f["rule_id"] == "convert.input_error"]
        self.assertEqual(errs, ["old.doc"])
        self.assertTrue((out_dir / "a.md").is_file())

    def test_llms_txt_written_even_when_all_failed(self):
        big = self._zip_bomb("b.xlsx")
        out_dir = self.tmp / "out"
        rc, _, _ = run_cli(str(big), "-d", str(out_dir), "--continue-on-error", "--max-unzipped-mb", "1", "--llms-txt")
        self.assertEqual(rc, 1)  # 続行した入力の誤りは、全件失敗でも変換失敗（1）として扱う
        self.assertTrue((out_dir / "llms.txt").is_file())
        self.assertTrue((out_dir / "llms-full.txt").is_file())

    def test_continue_on_error_all_ok_exits_zero(self):
        good = self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="a.xlsx")
        rc, _, _ = run_cli(str(good), "-d", str(self.tmp / "out"), "--continue-on-error")
        self.assertEqual(rc, 0)

    def test_argument_error_still_stops_immediately(self):
        good = self.book([["区分", "金額"], ["A", 1], ["B", 2]], name="a.xlsx")
        rc, _, err = run_cli(str(good), "--header-row", "x=y", "-d", str(self.tmp / "out"), "--continue-on-error")
        self.assertEqual(rc, 2)
        self.assertIn("--header-row", err)
        self.assertFalse((self.tmp / "out" / "a.md").exists())

    # -- 4. draw.io の zlib 展開の上限 -----------------------------------------
    def _compressed_drawio(self, xml: str) -> Path:
        import base64
        import urllib.parse
        import zlib

        c = zlib.compressobj(9, zlib.DEFLATED, -15)
        raw = c.compress(urllib.parse.quote(xml).encode()) + c.flush()
        path = self.tmp / "bomb.drawio"
        path.write_text(
            f'<mxfile><diagram name="p">{base64.b64encode(raw).decode()}</diagram></mxfile>', encoding="utf-8"
        )
        return path

    def test_drawio_inflate_limit(self):
        from devtools_common.cli import UsageError

        pad = "<!--" + "a" * (3 << 20) + "-->"
        path = self._compressed_drawio(f"<mxGraphModel>{pad}<root/></mxGraphModel>")
        with self.assertRaises(UsageError) as cm:
            self.convert(path, max_unzipped_mb=1)
        self.assertIn("上限", str(cm.exception))
        # 上限内なら読める
        r, _ = self.convert(path, max_unzipped_mb=8)
        self.assertEqual(r.source.name, "bomb.drawio")

    def test_drawio_inflate_limit_unit(self):
        import zlib

        from devtools_common.cli import UsageError

        c = zlib.compressobj(9, zlib.DEFLATED, -15)
        data = c.compress(b"x" * 5000) + c.flush()
        self.assertEqual(drawio_adapter._inflate(data, 5000, "n"), b"x" * 5000)
        with self.assertRaises(UsageError):
            drawio_adapter._inflate(data, 4999, "n")
