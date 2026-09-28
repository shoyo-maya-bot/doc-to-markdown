"""Word (.docx) → 中間表現（python-docx）.

- 本文の出現順（段落と表の混在）を保つ
- 見出し: 組み込みスタイル Title / Heading N、またはアウトラインレベル
- 箇条書き: 番号付け定義（numbering.xml）から番号付き / 記号付きを判定し、レベルを保持
- 太字・斜体は Markdown 記法に変換
- 段落内の画像は assets に書き出して参照
"""

from __future__ import annotations

import re
from pathlib import Path

import docx
from docx.oxml.ns import qn
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph

from ..assets import AssetWriter
from ..model import Document, Heading, Image, ListItem, Paragraph, Table

_HEADING_STYLE = re.compile(r"^(?:heading|見出し)\s*(\d)$", re.I)
_LIST_STYLE = re.compile(r"^List (Bullet|Number)(?: (\d))?$")  # 組み込みの箇条書きスタイル（"List Bullet 2" = 2 段目）


def _numbering_formats(document) -> dict[tuple[str, str], str]:
    """(numId, ilvl) → numFmt（bullet / decimal など）."""
    try:
        numbering = document.part.numbering_part.element
    except (KeyError, NotImplementedError, AttributeError):
        return {}
    abstract: dict[str, dict[str, str]] = {}
    for an in numbering.findall(qn("w:abstractNum")):
        levels = {}
        for lvl in an.findall(qn("w:lvl")):
            fmt = lvl.find(qn("w:numFmt"))
            levels[lvl.get(qn("w:ilvl"))] = fmt.get(qn("w:val")) if fmt is not None else "bullet"
        abstract[an.get(qn("w:abstractNumId"))] = levels
    result = {}
    for num in numbering.findall(qn("w:num")):
        ref = num.find(qn("w:abstractNumId"))
        if ref is None:
            continue
        for ilvl, fmt in abstract.get(ref.get(qn("w:val")), {}).items():
            result[(num.get(qn("w:numId")), ilvl)] = fmt
    return result


def _is_title(p: DocxParagraph) -> bool:
    return ((p.style.name if p.style is not None else "") or "").lower() == "title"


def _heading_level(p: DocxParagraph, offset: int = 0) -> int | None:
    """offset: 文書に Title スタイルがある場合は 1（Title を H1、Heading 1 を H2 にする）."""
    if _is_title(p):
        return 1
    name = (p.style.name if p.style is not None else "") or ""
    m = _HEADING_STYLE.match(name.strip())
    if m:
        return int(m.group(1)) + offset
    ppr = p._p.pPr
    if ppr is not None:
        ol = ppr.find(qn("w:outlineLvl"))
        if ol is not None:
            val = int(ol.get(qn("w:val")))
            if val < 9:
                return val + 1 + offset
    return None


def _list_info(p: DocxParagraph, formats) -> tuple[int, bool] | None:
    ppr = p._p.pPr
    numpr = ppr.find(qn("w:numPr")) if ppr is not None else None
    name = (p.style.name if p.style is not None else "") or ""
    style_list = _LIST_STYLE.match(name)
    style_level = int(style_list.group(2) or 1) - 1 if style_list else 0
    if numpr is None:
        if style_list:
            return style_level, style_list.group(1) == "Number"
        return None
    ilvl_el, numid_el = numpr.find(qn("w:ilvl")), numpr.find(qn("w:numId"))
    ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else str(style_level)
    num_id = numid_el.get(qn("w:val")) if numid_el is not None else None
    if num_id == "0":  # 番号付け解除
        return None
    fmt = formats.get((num_id, ilvl), "bullet" if "Bullet" in name else "decimal")
    return int(ilvl), fmt not in ("bullet", "none")


def _runs_to_markdown(p: DocxParagraph) -> str:
    parts: list[str] = []
    cur_fmt: tuple[bool, bool] | None = None
    buf = ""

    def flush():
        nonlocal buf
        if buf:
            b, i = cur_fmt or (False, False)
            core = buf.strip()
            if core and (b or i):
                mark = "***" if b and i else "**" if b else "*"
                lead, trail = buf[: len(buf) - len(buf.lstrip())], buf[len(buf.rstrip()) :]
                parts.append(f"{lead}{mark}{core}{mark}{trail}")
            else:
                parts.append(buf)
        buf = ""

    for run in p.runs:
        text = run.text
        if not text:
            continue
        fmt = (bool(run.bold), bool(run.italic))
        if fmt != cur_fmt:
            flush()
            cur_fmt = fmt
        buf += text
    flush()
    md = "".join(parts)
    return md if md.strip() else p.text


def _images_in(p: DocxParagraph, document, assets: AssetWriter) -> list[str]:
    paths = []
    for blip in p._p.iter(qn("a:blip")):
        rid = blip.get(qn("r:embed"))
        part = document.part.related_parts.get(rid) if rid else None
        if part is not None and hasattr(part, "blob"):
            paths.append(assets.write(part.blob, content_type=getattr(part, "content_type", "")))
    return paths


def _table_rows(t: DocxTable) -> list[list[str]]:
    rows = []
    for row in t.rows:
        cells, prev = [], None
        for cell in row.cells:
            # 横結合セルは python-docx では同じ要素が繰り返されるので、2 つ目以降は空にする
            cells.append("" if cell._tc is prev else cell.text)
            prev = cell._tc
        rows.append(cells)
    return rows


def convert(path: Path, assets: AssetWriter) -> Document:
    document = docx.Document(str(path))
    formats = _numbering_formats(document)
    doc = Document()
    st = doc.stats
    body = document.element.body
    offset = 1 if any(_is_title(p) for p in document.paragraphs) else 0
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = DocxParagraph(child, document)
            text = p.text
            level = _heading_level(p, offset)
            if text.strip():
                st.add_text(text)
                if level:
                    st.headings += 1
                    doc.blocks.append(Heading(level, text))
                else:
                    info = _list_info(p, formats)
                    if info:
                        st.list_items += 1
                        doc.blocks.append(ListItem(_runs_to_markdown(p), info[0], info[1]))
                    else:
                        doc.blocks.append(Paragraph(_runs_to_markdown(p)))
            for img in _images_in(p, document, assets):
                st.images += 1
                doc.blocks.append(Image(img, alt=Path(img).stem))
        elif child.tag == qn("w:tbl"):
            t = DocxTable(child, document)
            rows = _table_rows(t)
            st.tables += 1
            st.table_rows += sum(1 for r in rows if any(c.strip() for c in r))
            for r in rows:
                for c in r:
                    st.add_text(c)
            doc.blocks.append(Table(rows=rows, header=True))
    if document.sections and any(
        s.header and s.header.paragraphs and s.header.paragraphs[0].text for s in document.sections
    ):
        doc.notes.append("ヘッダ/フッタは本文に含めていません")
    return doc
