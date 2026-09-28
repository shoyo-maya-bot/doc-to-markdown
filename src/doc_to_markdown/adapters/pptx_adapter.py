"""PowerPoint (.pptx) → 中間表現（python-pptx）.

- スライドごとに H2（「N. タイトル」）
- 図形は上→下、左→右の順で読む（グループは展開）
- 本文プレースホルダの段落は箇条書き（段落レベルを保持）、その他のテキストボックスは段落
- 表・画像を抽出。発表者ノートは「ノート:」段落として末尾に付ける
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from ..assets import AssetWriter
from ..model import Document, Heading, Image, ListItem, Paragraph, Table

_TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE}
_BODY_TYPES = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.SUBTITLE}


def _ordered_shapes(shapes):
    def key(s):
        return (s.top or 0, s.left or 0)

    for s in sorted(shapes, key=key):
        if s.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _ordered_shapes(s.shapes)
        else:
            yield s


def _placeholder_type(shape):
    try:
        return shape.placeholder_format.type if shape.is_placeholder else None
    except ValueError:
        return None


def convert(path: Path, assets: AssetWriter) -> Document:
    prs = Presentation(str(path))
    doc = Document()
    st = doc.stats
    doc.blocks.append(Heading(1, path.stem, synthetic=True))
    for idx, slide in enumerate(prs.slides, 1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        if title:
            st.add_text(title)
        st.headings += 1
        doc.blocks.append(Heading(2, f"{idx}. {title}" if title else f"{idx}."))
        for shape in _ordered_shapes(slide.shapes):
            if title_shape is not None and shape.shape_id == title_shape.shape_id:
                continue
            ph = _placeholder_type(shape)
            if shape.has_text_frame:
                as_list = ph in _BODY_TYPES
                for para in shape.text_frame.paragraphs:
                    text = para.text
                    if not text.strip():
                        continue
                    st.add_text(text)
                    if as_list:
                        st.list_items += 1
                        doc.blocks.append(ListItem(text, para.level, False))
                    else:
                        doc.blocks.append(Paragraph(text))
            elif getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                st.tables += 1
                st.table_rows += sum(1 for r in rows if any(c.strip() for c in r))
                for r in rows:
                    for c in r:
                        st.add_text(c)
                doc.blocks.append(Table(rows=rows, header=True))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                img = shape.image
                st.images += 1
                rel = assets.write(img.blob, content_type=img.content_type, ext=img.ext)
                doc.blocks.append(Image(rel, alt=shape.name))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
            if notes:
                st.add_text(notes)
                doc.blocks.append(Paragraph("ノート: " + notes))
    doc.notes.append("SmartArt・グラフは対象外（必要なら --via-pdf を検討）")
    return doc
