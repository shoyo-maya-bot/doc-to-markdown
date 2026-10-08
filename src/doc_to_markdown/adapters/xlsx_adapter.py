"""Excel (.xlsx / .xlsm) → 中間表現（openpyxl）.

実務の Excel によくある形に対応する。

| パターン | 扱い |
|---|---|
| 表の上のタイトル行・作成者行 | 段落として表の前に出す（見出し行と取り違えない） |
| 見出し行が 1 行目でない | 自動検出（文字の入ったセルが多い最初の行）。`--header-row` で明示も可 |
| 2 段の見出し（「金額」の下に「税抜 / 税額 / 税込」） | 列名を「金額/税抜」のように連結 |
| セル結合 | 見出しは結合範囲の全列に同じ名前、本文の縦結合は値を繰り返す（検証しやすくするため） |
| 1 シートに複数の表 | 空行・空列で区切られた範囲ごとに別の表（`### 表N（A4:G11）`） |
| 表の下の注記（※…） | 表の後ろに段落として出す |
| 非表示の行・列・シート | 既定は出力して注記。`--skip-hidden` で除外 |
| 数式 | 保存されている計算結果を出す。結果が無いブックは警告（`--recalc` で LibreOffice 再計算） |
| セルの表示形式 | ゼロ埋め・桁区切り・小数桁・%・通貨・日付・和暦を表示形式どおりに出す（devtools_common.excel） |
| 経過時間（`[h]:mm` 等） | `36:00` のように表示どおりに出す（devtools_common.excel が未対応なのでここで補う） |
| 再現できない表示形式 | 書式なしの値を出し、`loss.number_format` で報告 |
| 年度・数値だけの見出し（`項目 / 2024 / 2025`） | 年度の並び・太字・文字列として入力された数値で見出し行と判定 |
| 1 列だけの表 | 見出しが太字、または本文がすべて数値で 3 行以上なら表 |
| 見出し行が見つからない表らしい範囲 | 段落として出し、`loss.table_as_paragraph` で報告（`--header-row` で指定できる） |
| 大きなブック | 20 万セル（`--max-cells`）を超えると省メモリ読み込み（1 シート = 1 表、結合の解決なし）。warning |
"""

from __future__ import annotations

import datetime as dt
import math
import re
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from devtools_common.excel import MergeMap, header_depth, header_names, looks_numeric, render_value
from devtools_common.tempfiles import secure_tempdir
from lxml import etree
from openpyxl.utils import get_column_letter

from ..assets import AssetWriter
from ..model import Block, Document, Heading, Options, Paragraph, StreamTable, Table, count_chars

NOTE_PREFIXES = ("※", "注", "備考", "＊", "*")


def format_value(v, number_format: str = "") -> str:
    return cell_text(v, number_format)[0]


def cell_text(v, number_format: str = "") -> tuple[str, bool]:
    """セルの値を表示形式どおりの文字列にする. 戻り値は (文字列, 再現できたか).

    経過時間（`[h]:mm` 等）は openpyxl が timedelta で返し、devtools_common.excel.render_value は
    `1 day, 12:00:00` になってしまうので、ここで表示どおりに出す。
    """
    if isinstance(v, dt.timedelta):
        text = render_elapsed(v, number_format)
        if text is not None:
            return text, True
        return render_elapsed(v, "[h]:mm:ss") or str(v), not number_format or number_format.lower() == "general"
    return render_value(v, number_format)


_ELAPSED_TOKEN = re.compile(r'\[(h+|m+|s+)\]|"[^"]*"|\\.|h+|m+|s+|\.0+|[:\s\-/.,]', re.I)


def render_elapsed(td: dt.timedelta, number_format: str) -> str | None:
    """経過時間を `[h]:mm` `[h]:mm:ss` `[mm]:ss` `[h]:mm:ss.00` などの表示形式で出す. 再現できなければ None."""
    fmt = number_format.split(";")[0]
    tokens: list[str] = []
    pos = 0
    while pos < len(fmt):
        m = _ELAPSED_TOKEN.match(fmt, pos)
        if not m:
            return None
        tokens.append(m.group(0))
        pos = m.end()
    top = next((t[1].lower() for t in tokens if t.startswith("[")), None)
    if top is None:
        return None
    frac = max((len(t) - 1 for t in tokens if t.startswith(".0")), default=0)
    unit = 10**frac
    total = round(abs(td.total_seconds()) * unit)  # 表示する最小の桁で丸める
    whole, part = divmod(total, unit)
    if top == "h":
        values = {"h": whole // 3600, "m": whole % 3600 // 60, "s": whole % 60}
    elif top == "m":
        values = {"h": 0, "m": whole // 60, "s": whole % 60}
    else:
        values = {"h": 0, "m": 0, "s": whole}
    out = []
    for t in tokens:
        low = t.lower()
        if low.startswith("["):
            out.append(str(values[low[1]]).zfill(len(low) - 2))
        elif low[0] in "hms":
            out.append(str(values[low[0]]).zfill(len(low)))
        elif low.startswith(".0"):
            out.append("." + str(part).zfill(frac))
        elif t.startswith('"'):
            out.append(t[1:-1])
        elif t.startswith("\\"):
            out.append(t[1:])
        else:
            out.append(t)
    return ("-" if td < dt.timedelta(0) else "") + "".join(out)


# ---------------------------------------------------------------------------
# シートの格子
# ---------------------------------------------------------------------------
class Grid:
    """1 シート分のセル値（文字列）と結合情報. 行・列は 1 始まり."""

    def __init__(self, ws, opts: Options):
        self.values: dict[tuple[int, int], str] = {}
        self.merges = MergeMap.from_worksheet(ws)  # 結合範囲内の全セル → 結合
        self.hidden_rows: set[int] = set()
        self.hidden_cols: set[int] = set()
        for r, dim in ws.row_dimensions.items():
            if dim.hidden:
                self.hidden_rows.add(r)
        for dim in ws.column_dimensions.values():
            if dim.hidden and dim.min and dim.max:
                self.hidden_cols.update(range(dim.min, dim.max + 1))
        # 再現できなかった表示形式 → そのセル（行, 列）の一覧
        self.unsupported_formats: dict[str, list[tuple[int, int]]] = {}
        # 見出し行の判定に使う書式・型（数値だけの見出しを、下の行との違いで見分ける）
        self.bold: set[tuple[int, int]] = set()
        self.typed_number: set[tuple[int, int]] = set()  # 数値として入力されたセル（文字列の "2024" は含まない）
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                text, faithful = cell_text(c.value, c.number_format)
                if not faithful:
                    self.unsupported_formats.setdefault(c.number_format, []).append((c.row, c.column))
                if text.strip():
                    self.values[(c.row, c.column)] = text
                    if c.font is not None and c.font.b:
                        self.bold.add((c.row, c.column))
                    if isinstance(c.value, (int, float)) and not isinstance(c.value, bool):
                        self.typed_number.add((c.row, c.column))
        self.skip_hidden = opts.skip_hidden

    def visible(self, r: int, c: int) -> bool:
        return not self.skip_hidden or (r not in self.hidden_rows and c not in self.hidden_cols)

    def output_text(self) -> str:
        """出力対象のセル値（非表示を除外する設定ならそれを除く）. ロス検知の独立カウントに使う."""
        return "".join(v for (r, c), v in self.values.items() if self.visible(r, c))

    def occupied(self, r: int, c: int) -> bool:
        m = self.merges.get(r, c)
        if m:
            return (m.r0, m.c0) in self.values
        return (r, c) in self.values

    def value(self, r: int, c: int, *, fill_merged: bool) -> str:
        m = self.merges.get(r, c)
        if m and fill_merged:
            return self.values.get((m.r0, m.c0), "")
        if m and (r, c) != (m.r0, m.c0):
            return ""
        return self.values.get((r, c), "")

    def bounds(self) -> tuple[int, int, int, int] | None:
        cells = list(self.values) + [(m.r1, m.c1) for m in self.merges if m.anchor in self.values]
        if not cells:
            return None
        return min(r for r, _ in cells), min(c for _, c in cells), max(r for r, _ in cells), max(c for _, c in cells)


def split_regions(g: Grid, r0: int, c0: int, r1: int, c1: int, depth: int = 0) -> list[tuple[int, int, int, int]]:
    """空行・空列で区切られた矩形に分割する（XY-cut）. 上→下、左→右の順."""

    def row_used(r: int) -> bool:
        return any(g.occupied(r, c) for c in range(c0, c1 + 1))

    def col_used(c: int, ra: int, rb: int) -> bool:
        return any(g.occupied(r, c) for r in range(ra, rb + 1))

    # 上下の空行を詰める
    while r0 <= r1 and not row_used(r0):
        r0 += 1
    while r1 >= r0 and not row_used(r1):
        r1 -= 1
    if r0 > r1:
        return []
    # 左右の空列を詰める
    while c0 <= c1 and not col_used(c0, r0, r1):
        c0 += 1
    while c1 >= c0 and not col_used(c1, r0, r1):
        c1 -= 1
    if depth > 6:
        return [(r0, c0, r1, c1)]
    # 空行で分割
    bands, start = [], r0
    for r in range(r0, r1 + 2):
        if r > r1 or not row_used(r):
            if start <= r - 1:
                bands.append((start, r - 1))
            start = r + 1
    if len(bands) > 1:
        return [x for a, b in bands for x in split_regions(g, a, c0, b, c1, depth + 1)]
    # 空列で分割
    cols, start = [], c0
    for c in range(c0, c1 + 2):
        if c > c1 or not col_used(c, r0, r1):
            if start <= c - 1:
                cols.append((start, c - 1))
            start = c + 1
    if len(cols) > 1:
        return [x for a, b in cols for x in split_regions(g, r0, a, r1, b, depth + 1)]
    return [(r0, c0, r1, c1)]


def _logical_keys(g: Grid, r: int, c0: int, c1: int) -> list[tuple[tuple[int, int], str]]:
    """行 r の、結合を 1 つと数えた (値のあるセル, 値) の一覧（空は除く）."""
    out, seen = [], set()
    for c in range(c0, c1 + 1):
        m = g.merges.get(r, c)
        key = (m.r0, m.c0) if m else (r, c)
        if key in seen:
            continue
        seen.add(key)
        v = g.values.get(key, "")
        if v and g.visible(*key):
            out.append((key, v))
    return out


def _logical_cells(g: Grid, r: int, c0: int, c1: int) -> list[str]:
    """行 r の、結合を 1 つと数えた値の一覧（空は除く）."""
    return [v for _, v in _logical_keys(g, r, c0, c1)]


_YEAR = re.compile(r"^(?:19|20|21)\d\d$")


def year_sequence(vals: list[str]) -> bool:
    """数値に見える値が 2 つ以上あり、すべて年（1900〜2199 の 4 桁）で、増える（減る）一方に並ぶ."""
    nums = [v.strip() for v in vals if looks_numeric(v)]
    if len(nums) < 2 or not all(_YEAR.match(v) for v in nums):
        return False
    years = [int(v) for v in nums]
    pairs = list(zip(years, years[1:], strict=False))
    return all(a < b for a, b in pairs) or all(a > b for a, b in pairs)


def mostly_text(vals: list[str]) -> bool:
    return sum(1 for v in vals if not looks_numeric(v)) >= len(vals) * 0.7


def _next_row(g: Grid, r: int, r1: int, c0: int, c1: int) -> int | None:
    for n in range(r + 1, r1 + 1):
        if _logical_keys(g, n, c0, c1):
            return n
    return None


def _differs_from_next(g: Grid, r: int, r1: int, c0: int, c1: int) -> bool:
    """数値を含む行が、すぐ下の行と見た目・型で区別できる（= 見出し行らしい）."""
    keys = _logical_keys(g, r, c0, c1)
    vals = [v for _, v in keys]
    n = _next_row(g, r, r1, c0, c1)
    if n is None:
        return False
    below = _logical_keys(g, n, c0, c1)
    if year_sequence(vals) and not year_sequence([v for _, v in below]):
        return True  # 項目 / 2024 / 2025 / 2026
    if all(k in g.bold for k, _ in keys) and not any(k in g.bold for k, _ in below):
        return True  # 見出しだけ太字
    # 見出しの数値は文字列として入力され、下の行の同じ列は数値
    numeric_cols = [k[1] for k, v in keys if looks_numeric(v)]
    below_typed = {k[1] for k in g.typed_number if k[0] == n}
    return bool(numeric_cols) and all((r, c) not in g.typed_number and c in below_typed for c in numeric_cols)


def detect_header_row(g: Grid, r0: int, c0: int, r1: int, c1: int) -> int | None:
    width = c1 - c0 + 1
    if width == 1:
        return detect_single_column_header(g, r0, c0, r1)
    need = max(2, math.ceil(width * 0.5))
    for r in range(r0, min(r1, r0 + 20) + 1):
        vals = _logical_cells(g, r, c0, c1)
        filled = sum(1 for c in range(c0, c1 + 1) if g.occupied(r, c))
        if filled >= need and len(vals) >= 2 and (mostly_text(vals) or _differs_from_next(g, r, r1, c0, c1)):
            return r
    return None


SINGLE_COLUMN_MIN_BODY = 3  # 1 列の表とみなす本文の最小行数（タイトル・作成者・注記の数行と取り違えない）
SINGLE_COLUMN_HEADER_MAX = 40  # 見出しにしては長すぎる文（段落）を除く


def detect_single_column_header(g: Grid, r0: int, c: int, r1: int) -> int | None:
    """1 列だけの範囲が表なら見出し行を返す.

    見出しがはっきりしている場合だけ表にする: 先頭行が短い文字列で、本文が 3 行以上あり、
    (a) 先頭行だけ太字、または (b) 本文がすべて数値。注記（※ 等）で始まる行があれば表にしない。
    """
    rows = [(r, g.values.get((r, c), "")) for r in range(r0, r1 + 1) if g.occupied(r, c) and g.visible(r, c)]
    if len(rows) < 1 + SINGLE_COLUMN_MIN_BODY:
        return None
    (h, head), body = rows[0], rows[1:]
    if not head or looks_numeric(head) or len(head) > SINGLE_COLUMN_HEADER_MAX or "\n" in head:
        return None
    if any(v.lstrip().startswith(NOTE_PREFIXES) for _, v in rows):
        return None
    bold_header = (h, c) in g.bold and not any((r, c) in g.bold for r, _ in body)
    numeric_body = all(looks_numeric(v) for _, v in body)
    return h if bold_header or numeric_body else None


def looks_like_table(g: Grid, r0: int, c0: int, r1: int, c1: int) -> bool:
    """見出し行が見つからなかった範囲が、表らしい（2 つ以上の値が並ぶ行が 2 行以上ある）か."""
    return sum(1 for r in range(r0, r1 + 1) if len(_logical_cells(g, r, c0, c1)) >= 2) >= 2


def _header_depth(g: Grid, h: int, c0: int, c1: int, r1: int) -> int:
    """見出しが何行あるか（結合が下の行まで伸びている、横結合の下に子見出しがある. devtools_common.excel）."""
    return header_depth(g.merges, lambda r, c: g.values.get((r, c), ""), h, first_col=c0, last_col=c1, last_row=r1)


def _header_names(g: Grid, h: int, depth: int, c0: int, c1: int) -> list[str]:
    return header_names(lambda r, c: g.value(r, c, fill_merged=True).replace("\n", " "), h, depth, range(c0, c1 + 1))


@dataclass
class SheetResult:
    blocks: list[Block]
    tables: int
    rows: int
    text: str
    filled: int
    # 表らしいのに見出し行が見つからず段落にした範囲（A1:C3 形式）
    flattened: list[str] = field(default_factory=list)
    # 表の末尾から外して段落にした注記行のセル（A1 形式）
    moved_notes: list[str] = field(default_factory=list)


def sheet_blocks(ws, g: Grid, opts: Options, sheet_title: str) -> SheetResult:
    blocks: list[Block] = []
    b = g.bounds()
    if b is None:
        return SheetResult([Paragraph("（空のシート）", synthetic=True)], 0, 0, "", 0)
    regions = split_regions(g, *b)
    forced = opts.header_rows.get(sheet_title, opts.header_rows.get("*"))
    table_regions = []
    for reg in regions:
        r0, c0, r1, c1 = reg
        if forced and r0 <= forced <= r1:
            h = forced
        elif forced:
            h = None
        else:
            h = detect_header_row(g, r0, c0, r1, c1)
        table_regions.append((reg, h))
    n_tables = sum(1 for _, h in table_regions if h is not None)

    count_tables = rows_total = filled = 0
    texts: list[str] = []
    flattened: list[str] = []
    moved_notes: list[str] = []
    for (r0, c0, r1, c1), h in table_regions:
        if h is None:  # 表でない範囲（タイトル・作成者・注記）
            if looks_like_table(g, r0, c0, r1, c1):
                flattened.append(f"{get_column_letter(c0)}{r0}:{get_column_letter(c1)}{r1}")
            for r in range(r0, r1 + 1):
                if opts.skip_hidden and r in g.hidden_rows:
                    continue
                vals = _logical_cells(g, r, c0, c1)
                if vals:
                    line = " / ".join(vals)
                    texts.append(line)
                    blocks.append(Paragraph(line))
            continue
        count_tables += 1
        # 見出し行より上（同じ範囲内）はタイトルとして段落に
        for r in range(r0, h):
            vals = _logical_cells(g, r, c0, c1)
            if vals:
                line = " / ".join(vals)
                texts.append(line)
                blocks.append(Paragraph(line))
        depth = _header_depth(g, h, c0, c1, r1)
        header = _header_names(g, h, depth, c0, c1)
        texts += [v for r in range(h, h + depth) for v in _logical_cells(g, r, c0, c1)]
        body_rows: list[list[str]] = []
        notes: list[str] = []
        last = r1
        # 表の末尾の注記行（1 セルだけで ※ 等で始まる）
        while last > h + depth - 1:
            vals = _logical_cells(g, last, c0, c1)
            if len(vals) == 1 and vals[0].lstrip().startswith(NOTE_PREFIXES):
                notes.insert(0, vals[0])
                col = next((c for c in range(c0, c1 + 1) if g.occupied(last, c)), c0)
                moved_notes.insert(0, f"{get_column_letter(col)}{last}")
                last -= 1
            else:
                break
        for r in range(h + depth, last + 1):
            if opts.skip_hidden and r in g.hidden_rows:
                continue
            row = []
            for c in range(c0, c1 + 1):
                m = g.merges.get(r, c)
                vertical = m is not None and m.r1 > m.r0 and m.c0 == c and (r, c) != (m.r0, m.c0)
                v = g.value(r, c, fill_merged=vertical)
                if vertical and v:
                    filled += 1
                row.append(v)
            if any(row):
                texts += [v for v in _logical_cells(g, r, c0, c1)]
                body_rows.append(row)
        if opts.skip_hidden:
            keep = [i for i, c in enumerate(range(c0, c1 + 1)) if c not in g.hidden_cols]
            header = [header[i] for i in keep]
            body_rows = [[row[i] for i in keep] for row in body_rows]
        if n_tables > 1:
            ref = f"{get_column_letter(c0)}{h}:{get_column_letter(c1)}{last}"
            blocks.append(Heading(3, f"表{count_tables}（{ref}）", synthetic=True))
        blocks.append(Table(rows=[header, *body_rows], header=True))
        rows_total += 1 + len(body_rows)
        for n in notes:
            texts.append(n)
            blocks.append(Paragraph(n))
    return SheetResult(blocks, count_tables, rows_total, "".join(texts), filled, flattened, moved_notes)


# ---------------------------------------------------------------------------
# 省メモリ（大きなブック）
# ---------------------------------------------------------------------------
def _row_texts(row) -> list[str]:
    """省メモリ読み込みの 1 行（ReadOnlyCell / EmptyCell）を表示形式どおりの文字列にする."""
    return [format_value(c.value, getattr(c, "number_format", None) or "") for c in row]


def _stream_sheet(ws, doc: Document) -> Iterator[Block]:
    it = ws.iter_rows()
    head: list[list[str]] = []
    for row in it:
        cells = _row_texts(row)
        if any(c.strip() for c in cells):
            head.append(cells)
            if len(head) >= 20:
                break
    if not head:
        yield Paragraph("（空のシート）", synthetic=True)
        return
    width = max(len(r) for r in head)
    header_idx = 0
    for i, r in enumerate(head):
        texts = [c for c in r if c.strip()]
        below = [c for c in head[i + 1] if c.strip()] if i + 1 < len(head) else []
        if len(texts) >= max(2, math.ceil(width * 0.5)) and (
            mostly_text(texts) or (year_sequence(texts) and not year_sequence(below))
        ):
            header_idx = i
            break
    for r in head[:header_idx]:
        line = " / ".join(c for c in r if c.strip())
        doc.stats.add_text(line)
        yield Paragraph(line)
    st = doc.stats
    st.tables += 1

    def rows():
        for r in head[header_idx:]:
            st.table_rows += 1
            st.add_text("".join(r))
            yield r
        for row in it:
            cells = _row_texts(row)
            if any(c.strip() for c in cells):
                st.table_rows += 1
                st.add_text("".join(cells))
                yield cells

    gen = rows()
    header = next(gen)
    yield StreamTable(header=header, rows=gen)


# ---------------------------------------------------------------------------
_SHEET = re.compile(r"xl/worksheets/sheet\d+\.xml$")
NS_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def scan_workbook(path: Path) -> tuple[int, int]:
    """(セル数, 計算結果が保存されていない数式セルの数) を、シートの XML を 1 回流し読みして数える.

    openpyxl で読むより速く、メモリも一定（大きなブックの判定に使う）。
    """
    cells = uncached = 0
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not _SHEET.match(name):
                continue
            with z.open(name) as f:
                for _, row in etree.iterparse(
                    f, events=("end",), tag=f"{NS_MAIN}row", resolve_entities=False, no_network=True, huge_tree=True
                ):
                    for c in row.iterchildren(f"{NS_MAIN}c"):
                        cells += 1
                        if c.find(f"{NS_MAIN}f") is not None:
                            v = c.find(f"{NS_MAIN}v")
                            if v is None or v.text is None:
                                uncached += 1
                    # 読み終えた行を木から外してメモリを一定に保つ
                    row.clear()
                    parent = row.getparent()
                    while row.getprevious() is not None:
                        del parent[0]
    return cells, uncached


def _blocks(path: Path, doc: Document, opts: Options, cells: int) -> Iterator[Block]:
    st = doc.stats
    yield Heading(1, path.stem, synthetic=True)
    full = cells <= opts.xlsx_full_load_max_cells
    wb = openpyxl.load_workbook(str(path), read_only=not full, data_only=True)
    if not full:
        doc.losses.append(
            (
                "convert.large_workbook",
                f"大きなブック（{cells:,} セル。上限 {opts.xlsx_full_load_max_cells:,}）のため"
                "省メモリで読み込みました: "
                "セル結合・複数の表・非表示・注記の判定は行わず、1 シート = 1 表で出力しています",
                "結合や複数の表を正しく出すには --max-cells でこのブックのセル数より大きな値を指定してください"
                "（メモリはセル 100 万個で約 800MB 使います）",
            )
        )
    raw = 0
    hidden_rows = hidden_sheets = filled = 0
    lost_formats: dict[str, list[str]] = {}  # 再現できなかった表示形式 → セル（シート名!A1）
    flattened: list[str] = []  # 表らしいのに段落にした範囲（シート名!A1:C3）
    moved_notes: list[str] = []  # 表の末尾から外して段落にした注記行（シート名!A10）
    try:
        for ws in wb.worksheets:
            hidden = getattr(ws, "sheet_state", "visible") != "visible"
            if hidden:
                hidden_sheets += 1
                if opts.skip_hidden:
                    continue
            st.headings += 1
            yield Heading(2, ws.title + ("（非表示）" if hidden else ""))
            if not full:
                yield from _stream_sheet(ws, doc)
                continue
            g = Grid(ws, opts)
            raw += count_chars(g.output_text())
            if g.hidden_rows:  # 値のある行を 1 回だけ集めてから数える（行 × セルの総当たりにしない）
                rows_with_values = {r for r, _ in g.values}
                hidden_rows += len(g.hidden_rows & rows_with_values)
            for fmt, cells in g.unsupported_formats.items():
                lost_formats.setdefault(fmt, []).extend(
                    f"{ws.title}!{get_column_letter(c)}{r}" for r, c in cells if g.visible(r, c)
                )
            res = sheet_blocks(ws, g, opts, ws.title)
            st.tables += res.tables
            st.table_rows += res.rows
            st.add_text(res.text)
            filled += res.filled
            flattened += [f"{ws.title}!{ref}" for ref in res.flattened]
            moved_notes += [f"{ws.title}!{ref}" for ref in res.moved_notes]
            yield from res.blocks
    finally:
        wb.close()
    if full:
        doc.raw_text_chars = raw
    if hidden_sheets:
        doc.notes.append(f"非表示のシート {hidden_sheets} 枚" + ("を除外" if opts.skip_hidden else "を含めて出力"))
    if hidden_rows and not opts.skip_hidden:
        doc.warnings.append(f"非表示の行 {hidden_rows} 行を含めて出力しています（除外するには --skip-hidden）")
    if moved_notes:
        doc.warnings.append(
            f"表の末尾の注記らしい行 {len(moved_notes)} 行（{'・'.join(NOTE_PREFIXES)} で始まる 1 セルだけの行）を"
            "表から外し、表の後に段落で出力しました: "
            + ", ".join(moved_notes[:5])
            + (f" ほか {len(moved_notes) - 5} 行" if len(moved_notes) > 5 else "")
            + "（表のデータなら --header-row で範囲を確かめるか、元の Excel で行頭の記号を見直してください）"
        )
    if filled:
        doc.notes.append(f"縦に結合されたセル {filled} 個は上の値を繰り返して出力")
    if flattened:
        doc.losses.append(
            (
                "loss.table_as_paragraph",
                f"見出し行が見つからないため、表らしい範囲 {len(flattened)} か所を"
                "段落（値を / で区切った行）で出力しました: "
                + ", ".join(flattened[:5])
                + (f" ほか {len(flattened) - 5} か所" if len(flattened) > 5 else ""),
                "表として出すには --header-row で見出し行を指定してください"
                "（例: --header-row 3 または --header-row シート名=3）",
            )
        )
    for fmt, refs in lost_formats.items():
        if refs:
            doc.losses.append(("loss.number_format", number_format_message(fmt, refs), NUMBER_FORMAT_SUGGESTION))


NUMBER_FORMAT_SUGGESTION = "元の Excel で該当セルの表示を確認し、必要なら Markdown を手で直してください"


def number_format_message(fmt: str, refs: list[str], shown: int = 5) -> str:
    cells = ", ".join(refs[:shown]) + (f" ほか {len(refs) - shown} セル" if len(refs) > shown else "")
    return f"表示形式 `{fmt}` を再現できないため、{len(refs)} セルを書式なしの値で出力しました: {cells}"


def convert(path: Path, assets: AssetWriter, opts: Options | None = None) -> Document:
    opts = opts or Options()
    doc = Document()
    cells, uncached = scan_workbook(path)
    if uncached and opts.recalc:
        doc.notes.append(f"数式の計算結果が無いセル {uncached} 個を LibreOffice で再計算して出力")
        doc.blocks = _recalc_blocks(path, doc, opts, cells)
        return doc
    if uncached:
        doc.warnings.append(
            f"数式の計算結果が保存されていないセルが {uncached} 個あり、空欄で出力しています"
            "（Excel で開いて保存し直すか、--recalc で LibreOffice に再計算させてください）"
        )
    doc.blocks = _blocks(path, doc, opts, cells)
    return doc


def _recalc_blocks(path: Path, doc: Document, opts: Options, cells: int) -> Iterator[Block]:
    from ..office import recalc_xlsx

    with secure_tempdir("doc2md-recalc-") as tmp:
        fixed = recalc_xlsx(path, tmp)
        for b in _blocks(fixed, doc, opts, cells):
            if isinstance(b, Heading) and b.synthetic and b.level == 1:
                b = Heading(1, path.stem, synthetic=True)
            yield b
